from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from datetime import datetime

def get_times(file):
    times = []
    with open(file) as f:
        startchunk = False
        f = list(f)
        for i, line in enumerate(f):
            lin = line.split()
            if not startchunk:
                time1 = lin[0]
                startchunk = True
            if lin == []:
                # the previous line
                time2 = f[i-1].split()[0]
                times.append([time1, time2])
                startchunk = False

    return times

def modify_excel_file(location_file):
    times = get_times(location_file)
    wb = load_workbook(filename="output.xlsx")
    ws = wb.active

    redFill = PatternFill(start_color='FFFF0000', end_color='FFFF0000', fill_type='solid')

    for cell in ws['A']:
        for time in times:
            time1 = datetime.strptime(time[0], '%H:%M:%S.%f')
            time2 = datetime.strptime(time[1], '%H:%M:%S.%f')
            if(type(cell.value) != str):
                if time1.time() <= cell.value <= time2.time():
                    cell.fill = redFill

    wb.save("processed.xlsx")

def func(location_file, data_file):
    times = get_times(location_file)
    with open(data_file) as f:
        f = list(f)
        for i, line in enumerate(f):
            for time in times:
                time1 = datetime.strptime(time[0], '%H:%M:%S.%f')
                time2 = datetime.strptime(time[1], '%H:%M:%S.%f')
                check = datetime.strptime(line.split()[0], '%H:%M:%S.%f')
                if time1.time() <= check.time() <= time2.time():
                    print(line, "A POINT")
                else:
                    print(line, "NOT A POINT")

if __name__ == "__main__":
    #print(get_times("outdoors.txt"))
    func("outdoors.txt", "raw.txt")
    modify_excel_file("outdoors.txt")